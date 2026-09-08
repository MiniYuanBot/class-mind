"""P1 输入网关（Input Gateway）。

职责：扫描输入目录、识别文件类型、提取课程元数据、识别课程类型、质量预检。
输入：input/ 目录下的 PDF / PPTX / DOCX / TXT / MD / 图片
输出：CourseMeta + FileManifest（文件清单与类型标注）+ course_type
"""

from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

from classmind.errors import InputError
from classmind.models import CourseMeta, CourseType

SLIDE_EXTS = {".pdf", ".pptx", ".ppt"}
TRANSCRIPT_EXTS = {".docx", ".txt", ".md"}
ASSET_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tiff"}

# 目录 / 文件名中暗示课程类型的启发词
_TYPE_HINTS: dict = {
    CourseType.THEORY: ["理论", "原理", "concept", "theory", "基础", "导论", "introduction", "overview"],
    CourseType.LAB: ["实验", "lab", "实验课", "实践", "实训"],
    CourseType.TUTORIAL: ["习题", "练习", "tutorial", "exercise", "作业", "例题", "辅导"],
    CourseType.SEMINAR: ["研讨", "seminar", "讲座", "前沿", "讨论", "论坛"],
}

# 章节号启发：第X讲 / 第X章 / lecture 5 / ch3 / L07
_CHAPTER_NO = re.compile(r"(?:第\s*([0-9一二三四五六七八九十百]+)\s*[讲章节课]|lecture\s*(\d+)|[Ll](?:esson)?\s*(\d+)|[Cc][Hh]\s*(\d+)|ch(?:apter)?\.?(\d+))", re.I)
_CN_NUM = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6, "七": 7, "八": 8, "九": 9, "十": 10}

COURSE_META_FILENAME = "meta.json"


def _cn_num_to_int(s: str) -> int:
    if s.isdigit():
        return int(s)
    if s in _CN_NUM:
        return _CN_NUM[s]
    if "十" in s:
        left = _CN_NUM.get(s[0], 1) if s[0] != "十" else 1
        right = _CN_NUM.get(s[-1], 0) if s[-1] != "十" else 0
        return left * 10 + right
    return 0


def _guess_chapter_no(path: Path) -> str:
    hay = path.stem
    m = _CHAPTER_NO.search(hay)
    if not m:
        # 也看父目录
        m = _CHAPTER_NO.search(path.parent.name)
    if m:
        for g in m.groups():
            if g:
                if g.isdigit():
                    return g
                return str(_cn_num_to_int(g))
    return ""


def _guess_chapter_title(path: Path) -> str:
    stem = path.stem
    for sep in ["-", "_", "—", " ", "："]:
        if sep in stem and any(k in stem for k in ["第", "lecture", "ch", "lesson", "讲", "章", "课"]):
            return stem.split(sep)[-1].strip()
    return ""


def _detect_type_from_text(*texts: str) -> Optional[CourseType]:
    blob = " ".join(texts).lower()
    score = {t: 0 for t in CourseType}
    for t, hints in _TYPE_HINTS.items():
        for h in hints:
            if h.lower() in blob:
                score[t] += 1
    best = max(score, key=lambda k: score[k])
    return best if score[best] > 0 else None


@dataclass
class Manifest:
    course_meta: CourseMeta = field(default_factory=CourseMeta)
    slides: list = field(default_factory=list)          # list[Path]
    transcripts: list = field(default_factory=list)     # list[Path]
    assets: list = field(default_factory=list)          # list[Path]
    ignored: list = field(default_factory=list)         # list[(Path, reason)]
    quality_issues: list = field(default_factory=list)

    @property
    def has_material(self) -> bool:
        return bool(self.slides or self.transcripts)


class InputGateway:
    """P1 输入网关。"""

    def __init__(self, input_dir: Path, overrides: Optional[dict] = None):
        self.input_dir = Path(input_dir)
        self.overrides = overrides or {}
        if not self.input_dir.exists():
            raise InputError(f"输入目录不存在: {self.input_dir}")

    # ------------------------------------------------------------------
    def scan(self) -> Manifest:
        files = sorted(
            (p for p in self.input_dir.rglob("*") if p.is_file()),
            key=lambda p: str(p).lower(),
        )
        manifest = Manifest(course_meta=CourseMeta())

        for path in files:
            ext = path.suffix.lower()
            name = path.name.lower()
            if name == COURSE_META_FILENAME:
                continue
            if ext in SLIDE_EXTS:
                manifest.slides.append(path)
            elif ext in TRANSCRIPT_EXTS:
                manifest.transcripts.append(path)
            elif ext in ASSET_EXTS:
                manifest.assets.append(path)
            else:
                manifest.ignored.append((path, f"不支持的文件类型 {ext or '(无扩展名)'}"))

        manifest.course_meta = self._resolve_meta(manifest)
        self._quality_precheck(manifest)
        return manifest

    # ------------------------------------------------------------------
    def _resolve_meta(self, manifest: Manifest) -> CourseMeta:
        ov = self.overrides
        meta_file = self.input_dir / COURSE_META_FILENAME
        file_meta: dict = {}
        if meta_file.exists():
            try:
                file_meta = json.loads(meta_file.read_text(encoding="utf-8"))
            except Exception as exc:  # noqa: BLE001
                manifest.quality_issues.append(f"[WARN] {COURSE_META_FILENAME} 无法解析，已忽略: {exc}")

        meta = CourseMeta()
        meta.source = "默认"
        name_src = ""
        explicit_type = False

        # 1) 显式 override（CLI 优先级最高）
        if ov.get("course_name"):
            meta.course_name = ov["course_name"]
            name_src = "CLI"
        if ov.get("instructor"):
            meta.instructor = ov["instructor"]
        if ov.get("chapter_no") is not None:
            meta.chapter_no = str(ov["chapter_no"])
        if ov.get("chapter_title"):
            meta.chapter_title = ov["chapter_title"]
        if ov.get("subject"):
            meta.subject = ov["subject"]
        if ov.get("course_code"):
            meta.course_code = str(ov["course_code"])
        if ov.get("lecture_date"):
            meta.lecture_date = str(ov["lecture_date"])
        if ov.get("file_stem"):
            meta.file_stem = str(ov["file_stem"])
        ctype = CourseType.parse(ov.get("course_type"))
        if ctype:
            meta.course_type = ctype
            explicit_type = True

        # 2) meta.json 补充缺失字段（文件名启发只做最后兜底）
        if file_meta:
            meta.source = f"meta.json + {name_src}" if name_src else "meta.json"
            for key, attr in [
                ("course", "course_name"), ("course_name", "course_name"), ("name", "course_name"),
                ("instructor", "instructor"), ("teacher", "instructor"),
                ("chapter", "chapter_no"), ("chapter_no", "chapter_no"), ("lecture", "chapter_no"),
                ("chapter_title", "chapter_title"), ("subject", "subject"), ("domain", "subject"),
                ("course_code", "course_code"), ("code", "course_code"),
                ("lecture_date", "lecture_date"), ("date", "lecture_date"),
                ("file_stem", "file_stem"), ("filename", "file_stem"),
            ]:
                cur = getattr(meta, attr)
                if (not cur or cur == "未命名课程") and file_meta.get(key):
                    setattr(meta, attr, str(file_meta[key]))
            ft = file_meta.get("course_type") or file_meta.get("type")
            parsed = CourseType.parse(ft)
            if parsed and not explicit_type:
                meta.course_type = parsed
                explicit_type = True

        # 3) 文件名 / 目录名启发填充（仅补空缺，不覆盖 meta.json / CLI）
        hint_names = [self.input_dir.name] + [p.stem for p in (manifest.slides + manifest.transcripts)]
        if not meta.course_name or meta.course_name == "未命名课程":
            candidates = [n for n in hint_names if n and not n.startswith(("第", "lecture", "ch"))]
            meta.course_name = candidates[0] if candidates else self.input_dir.name or "未命名课程"
            if "meta.json" not in meta.source:
                meta.source = (meta.source + " + 文件名启发").strip(" +")
        if not meta.chapter_no:
            for p in manifest.slides + manifest.transcripts:
                no = _guess_chapter_no(p)
                if no:
                    meta.chapter_no = no
                    break
        if not meta.chapter_title:
            for p in manifest.slides + manifest.transcripts:
                title = _guess_chapter_title(p)
                if title:
                    meta.chapter_title = title
                    break
        if not meta.subject:
            meta.subject = file_meta.get("subject", "") or ""

        # 4) 课程类型识别（仅当 meta.json 与 CLI 均未指定时采用文本启发）
        if not explicit_type:
            auto = _detect_type_from_text(*hint_names)
            if not auto and manifest.transcripts:
                sample = self._sample_text(manifest.transcripts[0], 1500)
                auto = _detect_type_from_text(sample)
            if auto:
                meta.course_type = auto
        return meta

    @staticmethod
    def _sample_text(path: Path, limit: int) -> str:
        try:
            if path.suffix.lower() == ".docx":
                text = _docx_to_text(path)
            else:
                text = path.read_text(encoding="utf-8", errors="ignore")
            return text[:limit]
        except Exception:  # noqa: BLE001
            return ""

    # ------------------------------------------------------------------
    def _quality_precheck(self, manifest: Manifest) -> None:
        if not manifest.has_material:
            manifest.quality_issues.append("[FAIL] 输入中未包含任何课件或讲述文本（至少需要一份）")
        if not manifest.slides:
            manifest.quality_issues.append("[WARN] 未发现课件（.pdf/.pptx），将仅依据讲述文本生成")
        if not manifest.transcripts:
            manifest.quality_issues.append("[WARN] 未发现讲述转录（.docx/.txt/.md），将仅依据课件生成")
        for p, reason in manifest.ignored:
            manifest.quality_issues.append(f"[WARN] 忽略 {p.name}: {reason}")


def _docx_to_text(path: Path) -> str:
    """stdlib 兜底 docx 文本抽取（P3 完整实现见 parsing/transcript_processor.py）。"""
    try:
        from docx import Document  # type: ignore

        doc = Document(str(path))
        return "\n".join(p.text for p in doc.paragraphs)
    except Exception:  # noqa: BLE001
        import re as _re
        import zipfile
        import xml.etree.ElementTree as ET

        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        try:
            with zipfile.ZipFile(path) as zf:
                xml = zf.read("word/document.xml")
            root = ET.fromstring(xml)
            paras = []
            for p in root.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"):
                texts = [t.text or "" for t in p.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t")]
                paras.append(_re.sub(r"\s+", " ", "".join(texts)).strip())
            return "\n".join(paras)
        except Exception:  # noqa: BLE001
            return ""
