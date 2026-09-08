"""P2 课件解析器（Slide2MD Parser）。

将 PDF / PPTX 课件解析为页级结构化 Markdown：
  - 页面标题  -> ## 页面标题
  - 要点列表  -> 无序列表（保留层级）
  - 公式      -> $$ ... $$
  - 表格      -> Markdown Table
  - 图片      -> 提取为独立资源 + ![描述](路径)
  - 页码锚点  -> <!-- slide: id=N, title="..." -->
输出：ParsedSlides（slides_md + slide_assets + slide_index）
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Optional

from classmind.errors import ParseError, UnsupportedFormatError
from classmind.models import ParsedSlides, Slide, SlideAsset, SlideIndexEntry
from classmind.util import tokenize_keywords

from .mdkit import latexify, page_anchor, table_to_md


class Slide2MDParser:
    """统一入口：按扩展名分发到 PDF / PPTX 解析器。"""

    def __init__(self, assets_dir: Path):
        self.assets_dir = Path(assets_dir)
        self.assets_dir.mkdir(parents=True, exist_ok=True)

    def parse(self, path: Path) -> ParsedSlides:
        ext = path.suffix.lower()
        if ext == ".pdf":
            return _parse_pdf(path, self.assets_dir)
        if ext in (".pptx", ".ppt"):
            return _parse_pptx(path, self.assets_dir)
        raise UnsupportedFormatError(f"课件格式不支持: {ext}（仅支持 .pdf / .pptx）")


# ---------------------------------------------------------------------------
# PDF（PyMuPDF）
# ---------------------------------------------------------------------------
def _parse_pdf(path: Path, assets_dir: Path) -> ParsedSlides:
    try:
        import fitz  # PyMuPDF
    except ImportError as exc:  # pragma: no cover
        raise UnsupportedFormatError("解析 PDF 需要 PyMuPDF，请执行: pip install classmind[pdf] 或 pip install PyMuPDF") from exc

    result = ParsedSlides(source_file=path.name)
    try:
        doc = fitz.open(path)
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"无法打开 PDF {path.name}: {exc}") from exc

    if doc.page_count == 0:
        result.warnings.append("[FAIL] PDF 页数为 0")
        doc.close()
        return result

    asset_counter = 0
    parts_md: list = []

    for pno in range(doc.page_count):
        page = doc[pno]
        slide_id = pno + 1
        text_dict = page.get_text("dict")
        blocks = _collect_text_blocks(text_dict)
        blocks = _drop_page_decorations(blocks, page)

        title, body_blocks = _pick_title(blocks, page)
        title = _clean_title(title)

        # ---- 表格（find_tables 可用时，带误判防护）----
        table_md_parts: list = []
        try:
            found = page.find_tables()
            for tab in found.tables:
                data = tab.extract()
                md = _table_md_if_valid(data, tab.bbox, page)
                if md:
                    table_md_parts.append(md)
        except Exception:  # noqa: BLE001
            pass

        # ---- 正文行（排除已入表的区域近似去重）----
        lines: list = []
        for blk in body_blocks:
            for ln in blk.get("lines", []):
                for span in ln.get("spans", []):
                    t = span.get("text", "")
                    if t.strip():
                        lines.append(t)
        if table_md_parts:
            lines = _subtract_table_cells(lines, table_md_parts)

        body = _body_to_md(lines)
        if table_md_parts:
            body = (body + "\n\n" + "\n\n".join(table_md_parts)).strip()

        # ---- 图片资源 ----
        image_refs = _extract_page_images(page, slide_id, assets_dir)
        asset_counter += len(image_refs)
        for ref in image_refs:
            desc, rel = ref
            body = body + f"\n\n![图 {desc}]({rel})"

        slide = Slide(
            index=slide_id,
            title=title,
            body_md=body,
            raw_text=_blocks_to_text(blocks),
            image_count=len(image_refs),
            table_count=len(table_md_parts),
            source_name=path.name,
            source_page=pno + 1,
        )
        result.slides.append(slide)
        parts_md.append(slide.markdown)

    doc.close()
    result.slides_md = "\n\n".join(parts_md)
    result.index = [
        SlideIndexEntry(
            page=s.index,
            title=s.title,
            keywords=tokenize_keywords(s.raw_text, 6) or tokenize_keywords(s.title, 6),
        )
        for s in result.slides
    ]
    if not any(s.raw_text.strip() or s.body_md.strip() for s in result.slides):
        result.warnings.append(f"[FAIL] PDF 无法提取文本（可能为扫描件/图片型 PDF），课件页数={doc.page_count}")
    return result


def _collect_text_blocks(text_dict: dict) -> list:
    """按阅读顺序（先 y 后 x）收集文本块。"""
    blocks = []
    for b in text_dict.get("blocks", []):
        if b.get("type", 0) != 0:
            continue  # 只取文本块（图片另由 image 处理）
        if not b.get("lines"):
            continue
        blocks.append(b)
    blocks.sort(key=lambda b: (round(b["bbox"][1], 1), round(b["bbox"][0], 1)))
    return blocks


def _drop_page_decorations(blocks: list, page) -> list:
    """去掉页眉页脚页码等装饰文本（如 Lec x.y / 日期 / © 页脚 / 纯页码）。"""
    ph = page.rect.height
    out = []
    for b in blocks:
        x0, y0, x1, y1 = b["bbox"]
        lines_out = []
        for l in b.get("lines", []):
            spans_text = " ".join(s.get("text", "") for s in l.get("spans", [])).strip()
            ly = (l["bbox"][1] + l["bbox"][3]) / 2
            if _is_decor_line(spans_text, ly, ph):
                continue
            if not spans_text:
                continue
            lines_out.append(l)
        if lines_out:
            # 重建块对象（只保留有效行）
            nb = dict(b)
            nb["lines"] = lines_out
            out.append(nb)
    return out


_DECOR_TOP = re.compile(r"^(Lec\s*[\d\.]*\s*$|CS\s*162|http)", re.I)
_DECOR_FOOT = re.compile(r"(©|UCB|Kubiatowicz|Berkeley|Fall\s*\d{4})", re.I)
_DECOR_DATE = re.compile(r"^\d{1,2}/\d{1,2}/\d{2,4}$")


def _is_decor_line(text: str, y_center: float, page_h: float) -> bool:
    if not text:
        return True
    # 纯页码（中下位置数字）
    if text.isdigit() and y_center > page_h * 0.85:
        return True
    if _DECOR_FOOT.search(text) and len(text) < 90:
        return True
    if _DECOR_TOP.match(text) and len(text) < 30:
        return True
    if _DECOR_DATE.match(text):
        return True
    return False


def _pick_title(blocks: list, page):
    """标题启发：页面上半部中字号最大的非装饰行作为标题。

    返回 (title, body_blocks)；title 所在块从正文中剔除。
    """
    ph = page.rect.height
    best = None  # (size, -y, text, block_id)
    best_block = None
    for b in blocks:
        for l in b.get("lines", []):
            sizes = [s.get("size", 0) for s in l.get("spans", [])]
            text = " ".join(s.get("text", "") for s in l.get("spans", [])).strip()
            ly = (l["bbox"][1] + l["bbox"][3]) / 2
            if not text or ly > ph * 0.55 or len(text) > 130:
                continue
            if _is_decor_line(text, ly, ph):
                continue
            size = max(sizes) if sizes else 0
            if size < 13:
                continue  # 正文小字号不做标题
            cand = (size, -ly, text, id(b))
            if best is None or cand > best:
                best = cand
                best_block = b
    if best is None:
        # 兜底：第一块文本行
        if blocks:
            b = blocks[0]
            for l in b.get("lines", []):
                text = " ".join(s.get("text", "") for s in l.get("spans", [])).strip()
                if text:
                    return _clean_title(text), blocks[1:]
        return "", blocks
    title = best[2].replace("\n", " ").strip()
    rest = [b for b in blocks if b is not best_block]
    return _clean_title(title), rest


def _clean_title(title: str) -> str:
    """清理页面标题：去项目符号 / 箭头 / 冗余空白，截断。"""
    t = (title or "").strip().replace("\n", " ")
    # 项目符号与装饰前缀（含常见 Unicode 圆点/箭头）
    t = re.sub(r"^[\s•·‣▪◦○●►→➢»»\-–—*]+", "", t)
    # 去掉末尾悬空逗号/冒号/破折号（如 “The first magnetic core memory,”）
    t = re.sub(r"[\s,，;；:：\-–—]+$", "", t)
    t = re.sub(r"\s+", " ", t)
    return t[:100]


def _table_md_if_valid(data, bbox, page):
    """把候选表格转 Markdown；防止把布局/装饰误判为表格。"""
    from .mdkit import table_to_md

    if not data or len(data) < 2:
        return ""
    headers = data[0]
    if len([c for c in headers if str(c).strip()]) < 2:
        return ""
    rows = [r for r in data[1:] if any(str(c).strip() for c in r)]
    if len(rows) < 1:
        return ""
    # 顶部 10% 区域（页眉条）与内容过小者不作为表格
    if bbox[1] < page.rect.height * 0.10:
        return ""
    if len(" ".join(str(c) for r in data for c in r).strip()) < 12:
        return ""
    return table_to_md(headers, rows)


def _blocks_to_text(blocks: list) -> str:
    lines = []
    for b in blocks:
        for l in b.get("lines", []):
            lines.append(" ".join(s.get("text", "") for s in l.get("spans", [])).strip())
    return "\n".join(lines)


def _body_to_md(lines: list) -> str:
    from .mdkit import block_to_md

    return block_to_md("", lines).strip()


def _subtract_table_cells(lines: list, tables_md: list) -> list:
    """近似去重：删除与表格单元格文本高度重合的行。"""
    cell_texts = set()
    for t in tables_md:
        for row in t.splitlines():
            if row.startswith("|"):
                for cell in row.strip("|").split("|"):
                    cell_texts.add(cell.strip())
    out = []
    for ln in lines:
        s = ln.strip()
        if s and s in cell_texts:
            continue
        out.append(ln)
    return out


def _extract_page_images(page, slide_id: int, assets_dir: Path) -> list:
    """提取页内真实显示的位图 -> assets/slide_XX_figN.png；返回 [(描述, 相对路径)]。"""
    refs = []
    try:
        import fitz

        page_rect = page.rect
        seen_xref = set()
        counter = 0
        for img in page.get_images(full=True):
            xref = img[0]
            if xref in seen_xref:
                continue
            seen_xref.add(xref)
            rects = page.get_image_rects(xref)
            for r in rects:
                if r.width * r.height < 2500 or r.width < 8 or r.height < 8:
                    continue
                if r.width > page_rect.width * 0.98 or r.height > page_rect.height * 0.98:
                    continue  # 背景/水印
                counter += 1
                fname = f"slide_{slide_id:02d}_fig{counter}.png"
                try:
                    pix = fitz.Pixmap(page.doc, xref)
                    if pix.n - pix.alpha > 3:
                        pix = fitz.Pixmap(fitz.csRGB, pix)
                    pix.save(str(assets_dir / fname))
                    refs.append((f"slide {slide_id} 第{counter}张图", f"assets/{fname}"))
                except Exception:  # noqa: BLE001
                    continue
    except Exception:  # noqa: BLE001
        pass
    return refs


# ---------------------------------------------------------------------------
# PPTX（python-pptx）
# ---------------------------------------------------------------------------
def _parse_pptx(path: Path, assets_dir: Path) -> ParsedSlides:
    try:
        from pptx import Presentation  # type: ignore
    except ImportError as exc:  # pragma: no cover
        raise UnsupportedFormatError("解析 PPTX 需要 python-pptx，请执行: pip install classmind[pptx] 或 pip install python-pptx") from exc

    result = ParsedSlides(source_file=path.name)
    try:
        prs = Presentation(str(path))
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"无法打开 PPTX {path.name}: {exc}") from exc

    parts_md: list = []
    for idx, slide in enumerate(prs.slides, start=1):
        title, blocks, notes = _collect_slide_content(slide, idx, path, assets_dir)
        title = (title or f"第 {idx} 页").strip()
        body_lines: list = []

        # 步骤拆解（动画近似）：以“步骤/Step N”开头处拆为子块
        step_groups = _split_steps(blocks)
        if len(step_groups) > 1:
            step_md: list = []
            for step_no, group in step_groups:
                step_md.append(page_anchor(idx, title, step_no))
                step_md.append(_render_step_group(group))
            body = "\n\n".join(step_md)
        else:
            body = "\n\n".join(blocks)

        if notes:
            body += f"\n\n> **讲师备注**：{notes}"

        raw_text = _blocks_to_plain(blocks)
        slide_obj = Slide(
            index=idx,
            title=title,
            body_md=body,
            raw_text=raw_text,
            image_count=len(_images_in_blocks(blocks)),
            table_count=len([b for b in blocks if b.startswith("|")]),
            source_name=path.name,
            source_page=idx,
            notes=notes,
        )
        result.slides.append(slide_obj)
        parts_md.append(slide_obj.markdown)

    result.slides_md = "\n\n".join(parts_md)
    result.index = [
        SlideIndexEntry(
            page=s.index,
            title=s.title,
            keywords=tokenize_keywords(s.raw_text or s.title, 6),
        )
        for s in result.slides
    ]
    if not result.slides:
        result.warnings.append("[FAIL] PPTX 无任何幻灯片")
    return result


def _collect_slide_content(slide, slide_id: int, source: Path, assets_dir: Path):
    """提取一张幻灯片的 (标题候选, 正文 Markdown 块 list, 备注)。"""
    title_candidate: str = ""
    blocks: list = []
    image_blocks: list = []

    def visit(shapes, parent_top=None):
        nonlocal title_candidate
        for shape in sorted(shapes, key=lambda s: (s.top if s.top is not None else 0, s.left if s.left is not None else 0)):
            # 图片
            if shape.shape_type == 13 or getattr(shape, "shape_type", None) == 13:  # MSO_SHAPE_TYPE.PICTURE
                try:
                    img = shape.image
                    ext = img.ext or "png"
                    fname = f"slide_{slide_id:02d}_img_{_safe_ext(ext)}"
                    blob = img.blob
                    (assets_dir / fname).write_bytes(blob)
                    image_blocks.append(f"![slide {slide_id} 图片](assets/{fname})")
                except Exception:  # noqa: BLE001
                    pass
                continue
            # 组合形状
            if shape.shape_type == 6:  # GROUP
                visit(shape.shapes, shape.top)
                continue
            # 表格
            if shape.has_table:
                tbl = shape.table
                rows = [[cell.text for cell in row.cells] for row in tbl.rows]
                if rows:
                    blocks.append(table_to_md(rows[0], rows[1:]))
                continue
            # 图表
            if getattr(shape, "has_chart", False) and shape.has_chart:
                try:
                    ch = shape.chart
                    cats = [str(c) for c in (ch.plots[0].categories if ch.plots else [])]
                    series_rows = []
                    headers = ["类别"]
                    for s in ch.series:
                        headers.append(s.name or "系列")
                    for i, cat in enumerate(cats):
                        row = [cat]
                        for s in ch.series:
                            vals = list(s.values)
                            row.append(str(vals[i]) if i < len(vals) else "")
                        series_rows.append(row)
                    if series_rows:
                        blocks.append(table_to_md(headers, series_rows))
                except Exception:  # noqa: BLE001
                    pass
                continue
            # 文本框 / 标题
            if shape.has_text_frame:
                text = _shape_text(shape)
                if not text.strip():
                    continue
                is_title = False
                if shape == slide.shapes.title:
                    is_title = True
                elif not title_candidate:
                    # 最大字号近似判断（EMU：1 inch = 914400，取页面上部 ~2.2in 内的大字号文本）
                    max_size = max((r.font.size.pt if r.font.size else 0) for p in shape.text_frame.paragraphs for r in p.runs)
                    area_top = shape.top if shape.top is not None else 0
                    if max_size >= 22 and area_top < 2_000_000:
                        is_title = True
                if is_title and not title_candidate:
                    title_candidate = text.replace("\n", " ")
                elif not is_title:
                    blocks.append(text)

    try:
        visit(slide.shapes)
    except Exception as exc:  # noqa: BLE001
        blocks.append(f"_[本页部分形状解析失败: {exc}]_")

    notes = ""
    try:
        if slide.has_notes_slide:
            notes = (slide.notes_slide.notes_text_frame.text or "").strip()
    except Exception:  # noqa: BLE001
        pass

    ordered = image_blocks + blocks
    return title_candidate, ordered, notes


def _safe_ext(ext: str) -> str:
    e = (ext or "png").lower()
    return e if e in {"png", "jpg", "jpeg", "gif", "bmp", "tiff"} else "png"


def _shape_text(shape) -> str:
    lines: list = []
    for para in shape.text_frame.paragraphs:
        runs = [r.text for r in para.runs if r.text]
        if not runs and para.text:
            runs = [para.text]
        text = "".join(runs)
        level = para.level or 0
        s = text.strip()
        if not s:
            continue
        indent = "  " * min(level, 3)
        if s.startswith(("•", "·", "-", "–")):
            lines.append(f"{indent}- {s[1:].strip()}")
        elif level > 0:
            lines.append(f"{indent}- {s}")
        else:
            lines.append(s)
    return "\n".join(lines)


def _blocks_to_plain(blocks: list) -> str:
    return "\n".join(b for b in blocks if not b.startswith("!["))[:4000]


def _images_in_blocks(blocks: list) -> list:
    return [b for b in blocks if b.startswith("![")]


def _split_steps(blocks: list) -> list:
    """把以“步骤N / Step N”开头的行拆成多组（模拟动画拆解）。

    工作在“行”粒度：同一文本框内的多行步骤也会被拆开。
    返回 list[(step_no, [lines, ...])]；若无步骤头则返回 [(1, blocks)]。
    """
    import re

    step_re = re.compile(r"^\s*(步骤|step)\s*[:：]?\s*\d", re.I)

    lines: list = []
    for b in blocks:
        lines.extend(b.split("\n"))
    boundaries = [i for i, ln in enumerate(lines) if step_re.match(ln)]
    if not boundaries:
        return [(1, blocks)]

    groups: list = []
    start = 0
    for i, boundary in enumerate(boundaries):
        if i == 0 and boundary > 0:
            groups.append(lines[start:boundary])
        elif i > 0:
            groups.append(lines[start:boundary])
        start = boundary
        if i == len(boundaries) - 1:
            groups.append(lines[boundary:])
    groups = [g for g in groups if g]
    if len(groups) <= 1:
        return [(1, blocks)]
    return [(i + 1, g) for i, g in enumerate(groups)]


def _render_step_group(group: list) -> str:
    """渲染一个步骤组：步骤头加粗，其余内容原样保留。"""
    import re

    head_re = re.compile(r"^\s*(步骤|step)\s*(\d+)\s*[:：]?\s*(.*)$", re.I)
    rendered: list = []
    for ln in group:
        m = head_re.match(ln)
        if m:
            rendered.append(f"**步骤 {m.group(2)}**：{m.group(3)}".rstrip())
        else:
            rendered.append(ln)
    return "\n".join(rendered)
